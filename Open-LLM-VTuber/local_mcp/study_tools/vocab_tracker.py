from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


CURRENT_DIR = Path(__file__).resolve().parent
OPEN_LLM_ROOT = CURRENT_DIR.parent.parent
DEFAULT_VOCAB_FILE = Path(
    os.getenv(
        "KURO_STUDY_VOCAB_FILE",
        str(OPEN_LLM_ROOT / "private" / "study" / "vocab.xlsx"),
    )
)
DEFAULT_STATE_DIR = Path(os.getenv("KURO_STUDY_STATE_DIR", str(OPEN_LLM_ROOT / "private" / "study")))

INDEX_FILE = "vocab_index.json"
PROGRESS_FILE = "vocab_progress.json"
SNAPSHOT_FILE = "study_snapshot.json"
EVENTS_FILE = "vocab_events.jsonl"

EXPECTED_COLUMNS = {
    "level": "JLPT等級",
    "kanji": "VocabKanji",
    "pos": "VocabPoS",
    "reading": "VocabFurigana",
    "meaning": "VocabDefTC",
    "sentence": "SentKanji1",
    "sentence_meaning": "SentDefTC1",
}

DEFAULT_LEVEL_ORDER = ["N3", "N2", "N4-5", "N1", "N5", "N4"]


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def today_key() -> str:
    return date.today().isoformat()


def clean_text(value: Any) -> str:
    text = str(value or "").strip()
    return re.sub(r"\s+", " ", text)


def read_json(path: Path, fallback: Any) -> Any:
    if not path.exists():
        return fallback
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return fallback


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def append_jsonl(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False) + "\n")


def file_metadata(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path),
        "size": stat.st_size,
        "mtime": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
    }


def sha1_short(value: str, length: int = 10) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:length]


def make_word_id(word: dict[str, Any]) -> str:
    level = clean_text(word.get("level")) or "unknown"
    kanji = clean_text(word.get("kanji")) or clean_text(word.get("reading")) or "word"
    reading = clean_text(word.get("reading"))
    meaning = clean_text(word.get("meaning"))
    digest = sha1_short("|".join([level, kanji, reading, meaning]))
    label = re.sub(r"[:\s]+", "_", kanji)[:24]
    return f"{level}:{label}:{digest}"


def column_index(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha())
    result = 0
    for letter in letters:
        result = result * 26 + (ord(letter.upper()) - ord("A") + 1)
    return max(result - 1, 0)


def xml_text(element: ET.Element) -> str:
    return "".join(element.itertext())


def read_xlsx_rows_stdlib(path: Path, sheet_name: str | None = None) -> list[list[Any]]:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    }

    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("main:si", ns):
                shared_strings.append(xml_text(item))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.attrib.get("Id"): rel.attrib.get("Target", "")
            for rel in rels.findall("pkg:Relationship", ns)
        }

        selected_rel = ""
        for sheet in workbook.findall("main:sheets/main:sheet", ns):
            name = sheet.attrib.get("name", "")
            if sheet_name is None or name == sheet_name:
                selected_rel = sheet.attrib.get(f"{{{ns['rel']}}}id", "")
                break
        if not selected_rel:
            raise ValueError(f"Sheet not found: {sheet_name}")

        target = rel_targets.get(selected_rel, "")
        sheet_path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
        root = ET.fromstring(archive.read(sheet_path))
        rows: list[list[Any]] = []
        for row in root.findall("main:sheetData/main:row", ns):
            values: list[Any] = []
            for cell in row.findall("main:c", ns):
                idx = column_index(cell.attrib.get("r", "A1"))
                while len(values) <= idx:
                    values.append("")
                cell_type = cell.attrib.get("t", "")
                if cell_type == "inlineStr":
                    inline = cell.find("main:is", ns)
                    values[idx] = xml_text(inline) if inline is not None else ""
                    continue
                value_node = cell.find("main:v", ns)
                raw = value_node.text if value_node is not None and value_node.text is not None else ""
                if cell_type == "s":
                    try:
                        values[idx] = shared_strings[int(raw)]
                    except Exception:
                        values[idx] = ""
                else:
                    values[idx] = raw
            rows.append(values)
        return rows


def read_xlsx_rows(path: Path, sheet_name: str | None = None) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except ImportError:
        return read_xlsx_rows_stdlib(path, sheet_name=sheet_name)


def rows_to_words(rows: list[list[Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    if not rows:
        return [], []

    headers = [clean_text(value) for value in rows[0]]
    header_lookup = {name: index for index, name in enumerate(headers) if name}
    missing = [column for column in EXPECTED_COLUMNS.values() if column not in header_lookup]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    words: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows[1:], start=2):
        def get(column_key: str) -> str:
            index = header_lookup[EXPECTED_COLUMNS[column_key]]
            return clean_text(row[index] if index < len(row) else "")

        kanji = get("kanji")
        reading = get("reading")
        meaning = get("meaning")
        if not kanji and not reading and not meaning:
            continue

        word = {
            "level": get("level"),
            "kanji": kanji,
            "pos": get("pos"),
            "reading": reading,
            "meaning": meaning,
            "sentence": get("sentence"),
            "sentenceMeaning": get("sentence_meaning"),
            "sourceRow": row_index,
        }
        word["word_id"] = make_word_id(word)
        words.append(word)

    return words, headers


def default_record(word: dict[str, Any]) -> dict[str, Any]:
    return {
        "word_id": word["word_id"],
        "status": "new",
        "ease": 2.5,
        "interval_days": 0,
        "seen_count": 0,
        "correct_count": 0,
        "wrong_count": 0,
        "streak": 0,
        "last_seen": "",
        "next_review": "",
        "suspended": False,
        "missingFromSource": False,
        "word": word,
    }


def merge_progress(existing: dict[str, Any], words: list[dict[str, Any]]) -> dict[str, Any]:
    records = existing.get("records") if isinstance(existing.get("records"), dict) else {}
    next_records: dict[str, Any] = {}
    for word in words:
        word_id = word["word_id"]
        current = records.get(word_id) if isinstance(records.get(word_id), dict) else {}
        merged = default_record(word)
        for key in (
            "status",
            "ease",
            "interval_days",
            "seen_count",
            "correct_count",
            "wrong_count",
            "streak",
            "last_seen",
            "next_review",
            "suspended",
        ):
            if key in current:
                merged[key] = current[key]
        merged["word"] = word
        next_records[word_id] = merged

    for word_id, current in records.items():
        if word_id not in next_records and isinstance(current, dict):
            missing = dict(current)
            missing["missingFromSource"] = True
            next_records[word_id] = missing

    return {
        "schemaVersion": 1,
        "updatedAt": now_iso(),
        "records": next_records,
    }


def due_date(record: dict[str, Any]) -> date | None:
    value = clean_text(record.get("next_review"))
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except Exception:
        return None


def is_due(record: dict[str, Any], day: date | None = None) -> bool:
    if record.get("suspended") or record.get("status") == "new":
        return False
    review_date = due_date(record)
    return bool(review_date and review_date <= (day or date.today()))


def level_rank(level: str, level_order: list[str]) -> int:
    try:
        return level_order.index(level)
    except ValueError:
        return len(level_order)


def record_sort_key(record: dict[str, Any], level_order: list[str]) -> tuple[Any, ...]:
    word = record.get("word") or {}
    return (
        level_rank(clean_text(word.get("level")), level_order),
        -(int(record.get("wrong_count") or 0)),
        int(word.get("sourceRow") or 999999),
        clean_text(word.get("kanji")),
    )


def public_word(record: dict[str, Any]) -> dict[str, Any]:
    word = record.get("word") or {}
    return {
        "word_id": record.get("word_id", ""),
        "level": word.get("level", ""),
        "kanji": word.get("kanji", ""),
        "reading": word.get("reading", ""),
        "meaning": word.get("meaning", ""),
        "pos": word.get("pos", ""),
        "sentence": word.get("sentence", ""),
        "sentenceMeaning": word.get("sentenceMeaning", ""),
        "status": record.get("status", "new"),
        "seen_count": int(record.get("seen_count") or 0),
        "correct_count": int(record.get("correct_count") or 0),
        "wrong_count": int(record.get("wrong_count") or 0),
        "next_review": record.get("next_review", ""),
    }


def choose_today_words(
    progress: dict[str, Any],
    *,
    review_count: int,
    new_count: int,
    focus_count: int,
    level_order: list[str],
) -> dict[str, list[dict[str, Any]]]:
    records = [
        record for record in progress.get("records", {}).values()
        if isinstance(record, dict) and not record.get("missingFromSource") and not record.get("suspended")
    ]
    today = date.today()

    due_records = [
        record for record in records
        if is_due(record, today) and record.get("status") != "mastered"
    ]
    due_records.sort(
        key=lambda record: (
            due_date(record) or today,
            -(int(record.get("wrong_count") or 0)),
            *record_sort_key(record, level_order),
        )
    )

    focus_records = [
        record for record in records
        if int(record.get("wrong_count") or 0) > 0 and record.get("status") != "mastered"
    ]
    focus_records.sort(key=lambda record: (-(int(record.get("wrong_count") or 0)), *record_sort_key(record, level_order)))

    used = {record.get("word_id") for record in due_records[:review_count]}
    focus_records = [record for record in focus_records if record.get("word_id") not in used]
    used.update(record.get("word_id") for record in focus_records[:focus_count])

    new_records = [record for record in records if record.get("status") == "new" and record.get("word_id") not in used]
    new_records.sort(key=lambda record: record_sort_key(record, level_order))

    return {
        "review": [public_word(record) for record in due_records[:review_count]],
        "focus": [public_word(record) for record in focus_records[:focus_count]],
        "new": [public_word(record) for record in new_records[:new_count]],
    }


def progress_stats(progress: dict[str, Any]) -> dict[str, Any]:
    records = [
        record for record in progress.get("records", {}).values()
        if isinstance(record, dict) and not record.get("missingFromSource")
    ]
    statuses = Counter(clean_text(record.get("status")) or "new" for record in records)
    levels = Counter(clean_text((record.get("word") or {}).get("level")) or "unknown" for record in records)
    due_count = sum(1 for record in records if is_due(record))
    wrong_count = sum(1 for record in records if int(record.get("wrong_count") or 0) > 0)
    return {
        "total": len(records),
        "statuses": dict(sorted(statuses.items())),
        "levels": dict(sorted(levels.items())),
        "due": due_count,
        "withMistakes": wrong_count,
    }


def build_snapshot(
    progress: dict[str, Any],
    *,
    review_count: int,
    new_count: int,
    focus_count: int,
    level_order: list[str],
) -> dict[str, Any]:
    chosen = choose_today_words(
        progress,
        review_count=review_count,
        new_count=new_count,
        focus_count=focus_count,
        level_order=level_order,
    )
    stats = progress_stats(progress)
    tasks = []
    if chosen["review"]:
        tasks.append({
            "id": "study:vocab-review",
            "source": "study",
            "kind": "vocab",
            "level": "must",
            "title": f"複習單字 {len(chosen['review'])} 個",
            "meta": "到期複習",
            "summary": "、".join(word["kanji"] for word in chosen["review"][:5]),
        })
    if chosen["new"]:
        tasks.append({
            "id": "study:vocab-new",
            "source": "study",
            "kind": "vocab",
            "level": "must",
            "title": f"新單字 {len(chosen['new'])} 個",
            "meta": "今日新增",
            "summary": "、".join(word["kanji"] for word in chosen["new"][:5]),
        })
    if chosen["focus"]:
        tasks.append({
            "id": "study:vocab-focus",
            "source": "study",
            "kind": "vocab",
            "level": "focus",
            "title": f"弱點單字 {len(chosen['focus'])} 個",
            "meta": "錯題優先",
            "summary": "、".join(word["kanji"] for word in chosen["focus"][:5]),
        })

    return {
        "schemaVersion": 1,
        "updatedAt": now_iso(),
        "source": "study",
        "studyType": "japanese",
        "vocab": {
            "stats": stats,
            "today": chosen,
        },
        "today": {
            "items": tasks,
        },
    }


def scan_vocab(args: argparse.Namespace) -> dict[str, Any]:
    vocab_file = Path(args.vocab_file)
    state_dir = Path(args.state_dir)
    if not vocab_file.exists():
        raise FileNotFoundError(f"Vocab file not found: {vocab_file}")

    rows = read_xlsx_rows(vocab_file, sheet_name=args.sheet)
    words, headers = rows_to_words(rows)
    existing_progress = read_json(state_dir / PROGRESS_FILE, {"records": {}})
    progress = merge_progress(existing_progress, words)

    index_payload = {
        "schemaVersion": 1,
        "updatedAt": now_iso(),
        "source": file_metadata(vocab_file),
        "sheet": args.sheet or "first",
        "columns": headers,
        "wordCount": len(words),
        "words": words,
    }
    snapshot = build_snapshot(
        progress,
        review_count=args.review_count,
        new_count=args.new_count,
        focus_count=args.focus_count,
        level_order=parse_levels(args.levels),
    )

    write_json(state_dir / INDEX_FILE, index_payload)
    write_json(state_dir / PROGRESS_FILE, progress)
    write_json(state_dir / SNAPSHOT_FILE, snapshot)

    return {
        "ok": True,
        "vocabFile": str(vocab_file),
        "stateDir": str(state_dir),
        "wordCount": len(words),
        "stats": snapshot["vocab"]["stats"],
        "todayCounts": {key: len(value) for key, value in snapshot["vocab"]["today"].items()},
    }


def parse_levels(value: str) -> list[str]:
    levels = [clean_text(item) for item in value.split(",") if clean_text(item)]
    return levels or DEFAULT_LEVEL_ORDER


def generate_today(args: argparse.Namespace) -> dict[str, Any]:
    state_dir = Path(args.state_dir)
    progress = read_json(state_dir / PROGRESS_FILE, None)
    if not isinstance(progress, dict):
        raise FileNotFoundError(f"Progress file not found. Run scan first: {state_dir / PROGRESS_FILE}")
    snapshot = build_snapshot(
        progress,
        review_count=args.review_count,
        new_count=args.new_count,
        focus_count=args.focus_count,
        level_order=parse_levels(args.levels),
    )
    write_json(state_dir / SNAPSHOT_FILE, snapshot)
    return {
        "ok": True,
        "stateDir": str(state_dir),
        "stats": snapshot["vocab"]["stats"],
        "todayCounts": {key: len(value) for key, value in snapshot["vocab"]["today"].items()},
    }


def next_review_after(days: int) -> str:
    return (date.today() + timedelta(days=max(0, days))).isoformat()


def mark_record(record: dict[str, Any], result: str) -> dict[str, Any]:
    next_record = dict(record)
    result = result.lower()

    if result == "suspend":
        next_record["suspended"] = True
        next_record["status"] = "suspended"
        return next_record
    if result == "unsuspend":
        next_record["suspended"] = False
        next_record["status"] = "review" if int(next_record.get("seen_count") or 0) else "new"
        return next_record
    if result == "reset":
        return default_record(record.get("word") or {})

    seen = int(next_record.get("seen_count") or 0) + 1
    correct = int(next_record.get("correct_count") or 0)
    wrong = int(next_record.get("wrong_count") or 0)
    streak = int(next_record.get("streak") or 0)
    ease = float(next_record.get("ease") or 2.5)
    interval = int(next_record.get("interval_days") or 0)

    if result in {"correct", "easy", "known"}:
        correct += 1
        streak += 1
        ease = min(3.2, ease + (0.18 if result == "easy" else 0.08))
        multiplier = ease * (1.35 if result in {"easy", "known"} else 1.0)
        interval = 1 if interval <= 0 else min(180, max(interval + 1, math.ceil(interval * multiplier)))
        status = "mastered" if streak >= 5 and interval >= 21 else "review"
        next_review = next_review_after(interval)
    elif result in {"wrong", "again"}:
        wrong += 1
        streak = 0
        ease = max(1.3, ease - 0.25)
        interval = 1
        status = "learning"
        next_review = next_review_after(1)
    elif result == "seen":
        status = "learning"
        interval = max(interval, 1)
        next_review = next_review_after(1)
    else:
        raise ValueError(f"Unsupported result: {result}")

    next_record.update({
        "status": status,
        "ease": round(ease, 2),
        "interval_days": interval,
        "seen_count": seen,
        "correct_count": correct,
        "wrong_count": wrong,
        "streak": streak,
        "last_seen": now_iso(),
        "next_review": next_review,
        "suspended": False,
    })
    return next_record


def mark_vocab(args: argparse.Namespace) -> dict[str, Any]:
    state_dir = Path(args.state_dir)
    progress = read_json(state_dir / PROGRESS_FILE, None)
    if not isinstance(progress, dict):
        raise FileNotFoundError(f"Progress file not found. Run scan first: {state_dir / PROGRESS_FILE}")

    records = progress.get("records")
    if not isinstance(records, dict) or args.word_id not in records:
        raise KeyError(f"Unknown word_id: {args.word_id}")

    before = records[args.word_id]
    after = mark_record(before, args.result)
    records[args.word_id] = after
    progress["updatedAt"] = now_iso()
    write_json(state_dir / PROGRESS_FILE, progress)
    append_jsonl(state_dir / EVENTS_FILE, {
        "time": now_iso(),
        "word_id": args.word_id,
        "result": args.result,
        "before": {
            "status": before.get("status"),
            "seen_count": before.get("seen_count"),
            "correct_count": before.get("correct_count"),
            "wrong_count": before.get("wrong_count"),
            "next_review": before.get("next_review"),
        },
        "after": {
            "status": after.get("status"),
            "seen_count": after.get("seen_count"),
            "correct_count": after.get("correct_count"),
            "wrong_count": after.get("wrong_count"),
            "next_review": after.get("next_review"),
        },
    })

    snapshot = build_snapshot(
        progress,
        review_count=args.review_count,
        new_count=args.new_count,
        focus_count=args.focus_count,
        level_order=parse_levels(args.levels),
    )
    write_json(state_dir / SNAPSHOT_FILE, snapshot)
    return {
        "ok": True,
        "word": public_word(after),
        "todayCounts": {key: len(value) for key, value in snapshot["vocab"]["today"].items()},
    }


def stats(args: argparse.Namespace) -> dict[str, Any]:
    state_dir = Path(args.state_dir)
    progress = read_json(state_dir / PROGRESS_FILE, None)
    if not isinstance(progress, dict):
        raise FileNotFoundError(f"Progress file not found. Run scan first: {state_dir / PROGRESS_FILE}")
    return {
        "ok": True,
        "stateDir": str(state_dir),
        "stats": progress_stats(progress),
    }


def print_json(payload: dict[str, Any]) -> None:
    sys.stdout.buffer.write(json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8") + b"\n")
    sys.stdout.buffer.flush()


def add_common_options(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--state-dir", default=str(DEFAULT_STATE_DIR), help="Directory for private progress files.")
    parser.add_argument("--review-count", type=int, default=20, help="Due review words to include in today's plan.")
    parser.add_argument("--new-count", type=int, default=10, help="New words to include in today's plan.")
    parser.add_argument("--focus-count", type=int, default=8, help="Mistake-prone words to include in today's plan.")
    parser.add_argument(
        "--levels",
        default=",".join(DEFAULT_LEVEL_ORDER),
        help="Comma-separated JLPT priority order for new words.",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Track Japanese vocabulary learning progress for Kuro.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    scan_parser = subparsers.add_parser("scan", help="Read the vocab workbook and update index/progress/snapshot.")
    add_common_options(scan_parser)
    scan_parser.add_argument("--vocab-file", default=str(DEFAULT_VOCAB_FILE), help="Source vocab xlsx file.")
    scan_parser.add_argument("--sheet", default="Vocab", help="Worksheet name to read.")

    today_parser = subparsers.add_parser("today", help="Regenerate today's study snapshot from progress.")
    add_common_options(today_parser)

    mark_parser = subparsers.add_parser("mark", help="Mark a word attempt result.")
    add_common_options(mark_parser)
    mark_parser.add_argument("--word-id", required=True, help="word_id from study_snapshot.json or vocab_progress.json.")
    mark_parser.add_argument(
        "--result",
        required=True,
        choices=["seen", "correct", "easy", "known", "wrong", "again", "suspend", "unsuspend", "reset"],
        help="Learning result to record.",
    )

    stats_parser = subparsers.add_parser("stats", help="Print progress stats.")
    stats_parser.add_argument("--state-dir", default=str(DEFAULT_STATE_DIR), help="Directory for private progress files.")

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        if args.command == "scan":
            print_json(scan_vocab(args))
        elif args.command == "today":
            print_json(generate_today(args))
        elif args.command == "mark":
            print_json(mark_vocab(args))
        elif args.command == "stats":
            print_json(stats(args))
        else:
            parser.error(f"Unsupported command: {args.command}")
    except Exception as exc:
        print_json({"ok": False, "error": str(exc), "errorType": exc.__class__.__name__})
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
